"""Fill the official SoftwareX Original Software Publication template with the manuscript text (content.py).

Input : paper/softwarex/template/softwarex-osp-template.docx (Elsevier template, v6 March 2026, not redistributed)
Output: paper/softwarex/pynhts_softwarex.docx and a word-count report.
Run   : python paper/softwarex/build_docx.py   (needs python-docx)
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import content as C  # noqa: E402

TEMPLATE = HERE / "template" / "softwarex-osp-template.docx"
OUT = HERE / "pynhts_softwarex.docx"
doc = docx.Document(str(TEMPLATE))
P = list(doc.paragraphs)  # indices as in the template


def drop(p) -> None:
    p._element.getparent().remove(p._element)


def fmt(run, size=11, bold=None, italic=False, mono=False) -> None:
    run.font.name = "Courier New" if mono else "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def set_text(p, text, size=11, bold=False, italic=False) -> None:
    for child in list(p._p):  # runs and hyperlinks of the template's placeholder text; keep the paragraph properties
        if not child.tag.endswith("}pPr"):
            p._p.remove(child)
    fmt(p.add_run(text), size=size, bold=bold, italic=italic)


def new_par(anchor, text="", *, size=11, bold=False, italic=False, mono=False, align=None, indent=None, after=6):
    p = anchor.insert_paragraph_before("", style="Normal")
    if text:
        fmt(p.add_run(text), size=size, bold=bold, italic=italic, mono=mono)
    p.paragraph_format.space_after = Pt(after)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    return p


WORDS: dict[str, int] = {}
FIGN = [0]
CITED: list[int] = []


def count(label: str, text: str) -> None:
    WORDS[label] = WORDS.get(label, 0) + len(re.findall(r"\S+", text))
    for grp in re.findall(r"\[(\d+(?:,\d+)*)\]", text):
        for n in grp.split(","):
            if int(n) not in CITED:
                CITED.append(int(n))


def emit(anchor, blocks, label: str) -> None:
    for b in blocks:
        kind = b[0]
        if kind == "p":
            new_par(anchor, b[1])
            count(label, b[1])
        elif kind == "h":
            new_par(anchor, b[1], bold=True, after=3)
            count(label, b[1])
        elif kind == "code":
            for line in b[1].split("\n"):
                new_par(anchor, line, size=8.5, mono=True, indent=0.2, after=0)
            new_par(anchor, "", after=4)
            count(label, b[1])
        elif kind == "fig":
            FIGN[0] += 1
            p = new_par(anchor, align="center", after=2)
            p.add_run().add_picture(str(HERE / "figures" / b[1]), width=Inches(5.6 if "fig1" in b[1] else 4.6))
            cap = new_par(anchor, align="center", after=8)
            fmt(cap.add_run(f"Fig. {FIGN[0]}. "), size=10, bold=True)
            fmt(cap.add_run(b[2]), size=10)
            count(label, f"Fig. {FIGN[0]}. {b[2]}")


# ---- front matter ------------------------------------------------------------------------------------
set_text(P[16], C.TITLE, size=14, bold=True)
set_text(P[18], "Seyed Saber Naseralavi", size=12, bold=True)
aff = new_par(P[19], "Department of Civil Engineering, Shahid Bahonar University of Kerman, Kerman 7616914111, Iran. "
              "ORCID 0000-0002-5392-910X. Corresponding author: [[AUTHOR-TODO: e-mail to be printed in the article]]", size=10)
set_text(P[21], C.ABSTRACT, size=11)
count("abstract", C.ABSTRACT)
set_text(P[25], C.KEYWORDS, size=11)
tbl = doc.tables[0]
for (key, val), row in zip(C.META, tbl.rows[1:], strict=True):
    assert row.cells[0].text.strip().startswith(key), (key, row.cells[0].text)
    cell = row.cells[2]
    for p in cell.paragraphs[1:]:
        drop(p)
    set_text(cell.paragraphs[0], val, size=10)

# ---- body --------------------------------------------------------------------------------------------
arch_main = [b for b in C.SECTIONS["architecture"]]
split = next(i for i, b in enumerate(arch_main) if b[0] == "h")
emit(P[38], C.SECTIONS["motivation"], "1 motivation")
emit(P[44], arch_main[:split], "2.1 architecture")
emit(P[46], arch_main[split + 1:], "2.2 functionalities")
emit(P[51], C.SECTIONS["examples"], "3 examples")
emit(P[61], C.SECTIONS["impact"], "4 impact")
emit(P[66], C.SECTIONS["conclusions"], "5 conclusions")

# ---- acknowledgements, declarations, references --------------------------------------------------------
set_text(P[69], C.ACK, size=11)
count("acknowledgements", C.ACK)
for head, text in (
    ("Declaration of generative AI and AI-assisted technologies in the manuscript preparation process", C.AI),
    ("CRediT authorship contribution statement", C.CREDIT),
    ("Funding", C.FUNDING),
    ("Declaration of competing interest", C.COMPETING),
    ("Data availability", C.DATA),
):
    new_par(P[70], head, bold=True, after=3)
    new_par(P[70], text)
    count("declarations", head + " " + text)
refs = {k: t for k, t in C.REFS}
for key, n in sorted(C.CITE_NUMBERS.items(), key=lambda kv: kv[1]):
    new_par(P[74], f"[{n}] {refs[key]}", size=10, after=3)

# ---- remove instructions and preamble ------------------------------------------------------------------
for i in [*range(0, 16), 17, 22, 29, 33, 34, 35, 36, 37, 41, 43, 45, 46, 50, 53, 54, 55, 56, 57, 58, 59, 60, 65, 72, 73, 76, 77, 78, 79]:
    drop(P[i])
doc.save(str(OUT))

body = sum(v for k, v in WORDS.items() if k != "declarations" and k != "acknowledgements")
print("words by part:", WORDS)
print("running text incl. abstract, captions, code; excluding declarations:", body, "| with declarations and acknowledgements:", sum(WORDS.values()))
print("figures:", FIGN[0], "| citation numbers in order of first appearance:", CITED)
assert CITED == sorted(CITED) == list(range(1, len(CITED) + 1)), "citations are not numbered in order of first appearance"
assert set(CITED) == set(C.CITE_NUMBERS.values()), "a reference is uncited or a citation has no reference"
